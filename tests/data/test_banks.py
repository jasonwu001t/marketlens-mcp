"""bank_fdic_financials, bank_fdic_industry and funds_mmf_statistics against
synthetic FDIC BankFind answers and a synthetic SEC Money Market Fund
Statistics page with two workbooks built here (fixtures/README)."""

from __future__ import annotations

import io
from datetime import UTC, datetime

import httpx
from data_harness import call, check_golden, context, fixture_json, refused

FDIC = r"api\.fdic\.gov/banks"
SEC = r"www\.sec\.gov"
LATE = datetime(2026, 10, 26, 20, 0, tzinfo=UTC)


def fdic_routes(upstream, monkeypatch=None):
    """The FDIC's three answers, and SEC's ticker list with the banks' holders.
    marketlens-data keeps SEC's list a day per process and on disk: with
    `monkeypatch`, both are read again, so this test's list is the one used."""
    if monkeypatch is not None:
        from omni.sources import sec as omni_sec

        monkeypatch.setattr(omni_sec, "_ticker_map_cache", None)
        monkeypatch.setattr(omni_sec, "_TICKER_TTL", 0)
    upstream.json_file(rf"{SEC}/files/company_tickers\.json", "fdic.company_tickers.json")
    upstream.json_file(rf"{SEC}/files/company_tickers_mf\.json", "sec.company_tickers_mf.json")
    upstream.json_file(rf"{FDIC}/institutions", "fdic.institutions.json")
    upstream.json_file(rf"{FDIC}/financials", "fdic.financials.json")
    upstream.json_file(rf"{FDIC}/summary", "fdic.summary.json")


def test_bank_fdic_financials_golden(tmp_path, upstream, keys, monkeypatch):
    fdic_routes(upstream, monkeypatch)
    out = call("bank_fdic_financials", context(tmp_path, now=LATE), tickers=["EXBK", "SELF", "TESTCO"])
    doc = check_golden(out, "bank_fdic_financials")
    assert [(r["ticker"], r["cert"], r["report_date"]) for r in doc["rows"]] == [
        ("EXBK", 99001, "2026-03-31"),
        ("EXBK", 99001, "2026-06-30"),
        ("EXBK", 99002, "2026-06-30"),
        ("SELF", 99003, "2026-03-31"),
        ("SELF", 99003, "2026-06-30"),
    ]
    first = doc["rows"][1]
    assert first["total_assets"] == 500_000_000_000.0, "the FDIC's thousands, in USD"
    assert first["htm_unrealized_gain"] == -5_000_000_000.0
    assert (first["roe"], first["cet1_ratio"], first["net_interest_margin"]) == (0.124, 0.132, 0.0295)
    assert first["knowledge_time"] == "2026-09-01T00:00:00Z", "the third month after the quarter"
    assert doc["rows"][0]["knowledge_time"] == "2026-06-01T00:00:00Z"
    selfheld = doc["rows"][3]
    assert selfheld["holding_company"] is None and selfheld["absent"]["holding_company"] == "not_applicable"
    assert selfheld["absent"]["cet1_ratio"] == "no_data"
    assert any("TESTCO" in n for n in doc["notes"]), "a company that owns no bank says so"
    q = [r for r in upstream.requests if r.url.path.endswith("/financials")][0]
    assert dict(q.url.params)["filters"] == "CERT:(99001 OR 99002 OR 99003)"
    early = call("bank_fdic_financials", context(tmp_path, now=LATE), tickers=["EXBK"], as_of="2026-08-31")
    assert [r.report_date.isoformat() for r in early.rows] == ["2026-03-31"]


def test_bank_fdic_industry_golden(tmp_path, upstream, keys):
    fdic_routes(upstream)
    out = call("bank_fdic_industry", context(tmp_path, now=LATE), start_year=2024)
    doc = check_golden(out, "bank_fdic_industry")
    assert [(r["year_end"], r["charter_class"], r["institutions"]) for r in doc["rows"]] == [
        ("2024-12-31", "CB", 3900),
        ("2024-12-31", "SI", 540),
        ("2025-12-31", "CB", 3800),
        ("2025-12-31", "SI", 520),
    ]
    assert doc["rows"][2]["total_assets"] == 24_000_000_000_000.0
    # the newest year at its index's build, an older one March 1 after it
    assert doc["rows"][2]["knowledge_time"] == "2026-09-16T13:42:38Z"
    assert doc["rows"][0]["knowledge_time"] == "2025-03-01T00:00:00Z"
    ca = call("bank_fdic_industry", context(tmp_path, now=LATE), start_year=2025, location="california")
    assert [(r.location, r.charter_class) for r in ca.rows] == [("California", "CB"), ("California", "SI")]
    err = refused("bank_fdic_industry", context(tmp_path, now=LATE), start_year=2026, end_year=2020)
    assert err.code == "invalid_arguments"


# --- money market funds ---------------------------------------------------------------------

TITLES = {
    "Tab.2.1": "MMF Net Assets ($Billions)",
    "Tab.2.6": "MMF Total Assets ($Billions)",
    "Tab.3.1": "Aggregate Gross Subscriptions, by MMF Category ($Billions)",
    "Tab.3.5": "Aggregate Gross Redemptions, by MMF Category ($Billions)",
    "Tab.3.9": "Aggregate Net Flows, by MMF Category ($Billions)",
    "Tab.4.3": "Asset-Weighted Seven Day Net Yield (Percent), by MMF Category",
    "Tab.5.1": "Asset-Weighted WAM, by MMF Category (Days)",
    "Tab.5.3": "Asset-Weighted WAL, by MMF Category (Days)",
    "Tab.6.1": "Aggregate MMF Daily Liquid Assets (Percent of Total Assets)",
    "Tab.6.3": "Aggregate MMF Weekly Liquid Assets (Percent of Total Assets)",
    "Tab.7.1": "Aggregate Value of MMF Investments, by Category ($Billions)",
    "Tab.7.4": "Aggregate Value of Government MMF Investments, by Category ($Billions)",
    "Tab.7.7": "Aggregate Value of Prime MMF Investments, by Category ($Billions)",
    "Tab.8.2": "Aggregate Value of MMF Repurchase Agreements with the Federal Reserve, by Category ($Billions)",
}


def workbook(months: list[datetime], net_assets: list[float], fed_repo: list[float]) -> bytes:
    """A synthetic workbook in the SEC's layout: a Notes sheet titling each
    table, each table a Category / item column and one column per month."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    notes = wb.create_sheet("Notes")
    for sheet, title in TITLES.items():
        notes.append([sheet, f"{title}\nSynthetic notes."])
    for sheet in TITLES:
        ws = wb.create_sheet(sheet)
        ws.append(["Category", "Fund Type", *months, None])
        if sheet == "Tab.2.1":
            ws.append(["Total", "Total", *net_assets])
        elif sheet == "Tab.8.2":
            ws.append(["Total", "Total", *fed_repo])
        elif sheet == "Tab.4.3":
            ws.append(["Government", "Government", *[3.5 + i / 10 for i in range(len(months))]])
        else:
            ws.append(["Total", "Total", *[1.0 for _ in months]])
        ws.append(["Return to Notes", None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


AUG = [datetime(2026, 7, 1), datetime(2026, 8, 1)]
SEP_ = [datetime(2026, 7, 1), datetime(2026, 8, 1), datetime(2026, 9, 1)]


def mmf_routes(upstream):
    upstream.text_file(
        rf"{SEC}/data-research/investment-management-data/money-market-fund-statistics",
        "sec.mmf_statistics.html",
    )
    books = {
        "08-2026": (workbook(AUG, [8400.0, 8440.0], [0.0, 0.0]), "Thu, 24 Sep 2026 17:54:51 GMT"),
        # September's workbook revises August's net assets
        "09-2026": (
            workbook(SEP_, [8400.0, 8439.5, 8384.1], [0.0, 0.0, 6.114]),
            "Fri, 23 Oct 2026 14:22:27 GMT",
        ),
    }

    def respond(request):
        month = request.url.path.rsplit("-data-", 1)[1].removesuffix(".xlsx")
        body, stamp = books[month]
        return httpx.Response(200, content=body, headers={"Last-Modified": stamp})

    upstream.add(rf"{SEC}/files/investment/mmf-statistics-supporting-data-\d\d-2026\.xlsx", respond)


def test_funds_mmf_statistics_golden(tmp_path, upstream, keys):
    mmf_routes(upstream)
    ctx = context(tmp_path, now=LATE)
    out = call("funds_mmf_statistics", ctx, start="2026-07-01", tables=["2.1", "8.2", "4.3"])
    doc = check_golden(out, "funds_mmf_statistics")
    net = [(r["month"], r["value"], r["unit"], r["file"]) for r in doc["rows"] if r["table"] == "2.1"]
    assert net == [
        ("2026-07-01", 8.4e12, "USD", "mmf-statistics-supporting-data-08-2026.xlsx"),
        ("2026-08-01", 8.4395e12, "USD", "mmf-statistics-supporting-data-09-2026.xlsx"),
        ("2026-09-01", 8.3841e12, "USD", "mmf-statistics-supporting-data-09-2026.xlsx"),
    ], "August as September's workbook revised it"
    fed = [r for r in doc["rows"] if r["table"] == "8.2" and r["month"] == "2026-09-01"]
    assert fed[0]["value"] == 6.114e9 and fed[0]["knowledge_time"] == "2026-10-24T03:59:59.999999Z"
    yld = [r for r in doc["rows"] if r["table"] == "4.3"]
    assert yld[0]["unit"] == "fraction" and yld[0]["value"] == 0.035
    # August as first published, before September's workbook
    before = call("funds_mmf_statistics", ctx, start="2026-07-01", tables=["2.1"], as_of="2026-10-01")
    assert [(r.month.isoformat(), r.value) for r in before.rows] == [
        ("2026-07-01", 8.4e12),
        ("2026-08-01", 8.44e12),
    ]
    err = refused("funds_mmf_statistics", ctx, tables=["9.9"])
    assert err.code == "invalid_arguments"


def test_the_synthetic_answers_are_marked():
    for name in ("fdic.institutions.json", "fdic.financials.json", "fdic.summary.json"):
        assert fixture_json(name)["_synthetic"] is True
